"""Anrufliste bauen: Recherche-JSONs eines Auftrags -> Excel im Anruflisten-Format.

Aufrufe:
  python3 build.py neu "<Titel>"        legt ~/Anruflisten/<datum>-<titel>/ an, gibt den Pfad aus
  python3 build.py ausschluss <ordner>  schreibt <ordner>/ausschluss_domains.txt (alle früher gelieferten Firmen)
  python3 build.py excel <ordner>       baut <ordner>/<Titel>-Anrufliste-<datum>.xlsx aus <ordner>/data/*.json
"""
import datetime, glob, json, os, re, subprocess, sys

# openpyxl bei Bedarf in einer eigenen venv nachinstallieren (System-Python hat es oft nicht)
try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    venv = os.path.expanduser("~/.cache/anrufliste-venv")
    py = os.path.join(venv, "Scripts", "python.exe") if os.name == "nt" else os.path.join(venv, "bin", "python3")
    if os.path.realpath(sys.executable) == os.path.realpath(py):
        sys.exit(f"openpyxl ließ sich nicht installieren (Ordner {venv} löschen und erneut versuchen)")
    try:
        if not os.path.exists(py):
            subprocess.check_call([sys.executable, "-m", "venv", venv])
        if subprocess.call([py, "-c", "import openpyxl"], stderr=subprocess.DEVNULL) != 0:
            subprocess.check_call([py, "-m", "pip", "install", "-q", "--disable-pip-version-check", "openpyxl"])
    except (subprocess.CalledProcessError, OSError) as e:
        sys.exit(f"openpyxl ließ sich nicht installieren ({e}); Internet prüfen und erneut versuchen")
    if os.name == "nt":  # execv ersetzt unter Windows den Prozess nicht sauber
        sys.exit(subprocess.call([py] + sys.argv))
    os.execv(py, [py] + sys.argv)
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

def _repo_root():
    """Liegt der Skill in <repo>/.claude/skills/anrufliste (Claude Code im Browser), gehören die Listen ins Repo."""
    hier = os.path.dirname(os.path.realpath(__file__))
    repo = os.path.dirname(os.path.dirname(os.path.dirname(hier)))
    if os.path.basename(os.path.dirname(hier)) == "skills" and repo != os.path.realpath(os.path.expanduser("~")) \
            and os.path.exists(os.path.join(repo, ".git")):
        return repo
    return None


REPO = _repo_root()
ROOT = os.path.expanduser(os.environ.get("ANRUFLISTEN_DIR") or (os.path.join(REPO, "Anruflisten") if REPO else "~/Anruflisten"))
STRICH = "–"
STATUS = ["Termin", "Rückruf vereinbart", "Chef nicht erreicht", "kein Bedarf (6 Mon.)", "gesperrt"]
RANG = {"hoch": 0, "mittel": 1, "niedrig": 2}

HEAD = PatternFill("solid", fgColor="1F4E78")
STUFE_FILL = {"A": PatternFill("solid", fgColor="C6EFCE"), "B": PatternFill("solid", fgColor="FFEB9C"), "C": PatternFill("solid", fgColor="F2F2F2")}
GRUEN = STUFE_FILL["A"]
ORANGE = PatternFill("solid", fgColor="F8CBAD")
ROT = PatternFill("solid", fgColor="FFC7CE")
WRAP = Alignment(wrap_text=True, vertical="top")
LINK = Font(color="0563C1", underline="single")

FELDER = ("firma", "ort", "kreis", "bundesland", "ansprechpartner", "rolle", "tel_ansprechpartner", "tel_zentrale", "email",
          "website", "mitarbeiter", "mitarbeiter_sicher", "widerspruch", "erreichbarkeit", "kaufsignal", "signal_staerke",
          "eignung", "notiz", "verbund", "konkurrenz")


def sauber(v):
    """Steuerzeichen entfernen (openpyxl lehnt sie ab) und auf Excels Zellgrenze kürzen."""
    v = ILLEGAL_CHARACTERS_RE.sub("", str(v))
    return v if len(v) <= 32000 else v[:32000] + " …"


def leer(v):
    if v is None:
        return STRICH
    if isinstance(v, list):
        v = "; ".join(str(x) for x in v if x)
    elif isinstance(v, dict):
        v = "; ".join(f"{k}: {x}" for k, x in v.items() if x)
    v = sauber(v).strip()
    return v if v and v not in ("-", "—", "k.A.", "k. A.", "None", "null", "n/a") else STRICH


def als_liste(v):
    """Quellen in eine Liste von Texten bringen, egal ob Agent Text, Liste oder Objekt liefert."""
    if v is None or v == "":
        return []
    if isinstance(v, dict):
        return [f"{k}: {x}" for k, x in v.items()]
    if isinstance(v, list):
        return [leer(x) for x in v if x]
    return [leer(v)]


def domain(url):
    w = str(url or "").strip().lower().split(" ")[0]
    w = re.sub(r"^[a-z]+://", "", w)
    w = re.split(r"[/?#]", w)[0].split("@")[-1].split(":")[0]
    w = re.sub(r"^www\d?\.", "", w).rstrip(".")
    return w if re.fullmatch(r"[a-z0-9äöü.-]+\.[a-z]{2,}", w) else ""


def json_liste(p):
    """Firmenliste aus einer JSON-Datei; None bei kaputtem oder falsch aufgebautem Inhalt."""
    try:
        daten = json.load(open(p, encoding="utf-8"))
    except (ValueError, OSError) as e:
        print("FEHLER, Datei übersprungen:", p, e)
        return None
    if isinstance(daten, dict):
        daten = daten.get("firmen", [daten])
    if not isinstance(daten, list):
        print("FEHLER, keine Firmenliste, übersprungen:", p)
        return None
    return [r for r in daten if isinstance(r, dict)]


def slug(s):
    s = s.strip().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("Ä", "Ae").replace("Ö", "Oe").replace("Ü", "Ue").replace("ß", "ss")
    return re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-")[:60] or "Liste"


def wort(s):
    """Erstes Wort klein (für hoch/mittel/niedrig)."""
    return (re.findall(r"[a-zäöü]+", leer(s).lower()) or [""])[0]


def ma_zahl(s):
    nums = [int(n.replace(".", "")) for n in re.findall(r"\d[\d.]*", s or "") if n.replace(".", "").isdigit()]
    nums = [n for n in nums if n < 100000]
    return nums[0] if nums else None


def eingeschraenkt(e):
    return e.lower().startswith("eingeschränkt")


# ---------- früher gelieferte Firmen ----------

def alte_domains(ausser=None):
    """Domains aus allen früheren Aufträgen (data/*.json) und aus Excel-Dateien in ~/Anruflisten/alt/."""
    doms = set()
    for p in glob.glob(os.path.join(ROOT, "*", "data", "*.json")):
        if ausser and os.path.realpath(p).startswith(os.path.realpath(ausser) + os.sep):
            continue
        for r in json_liste(p) or []:
            if domain(r.get("website")):
                doms.add(domain(r.get("website")))
    for p in glob.glob(os.path.join(ROOT, "alt", "*.xlsx")):
        try:
            wb = load_workbook(p, read_only=True)
        except Exception as e:
            print("alte Liste nicht lesbar, übersprungen:", p, e)
            continue
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            kopf_ = [str(c or "").lower() for c in next(rows, [])]
            if "website" not in kopf_:
                continue
            i = kopf_.index("website")
            for r in rows:
                if r and i < len(r) and domain(r[i]):
                    doms.add(domain(r[i]))
    for p in glob.glob(os.path.join(ROOT, "alt", "*.txt")):
        doms |= {domain(x) or x.strip().lower() for x in open(p, encoding="utf-8").read().split() if x.strip()}
    if REPO and not os.environ.get("ANRUFLISTEN_DIR"):
        doms |= git_domains(ausser)
    return doms


def git_domains(ausser=None):
    """Im Browser liegt jede Liste auf eigenem Branch: Firmen aus Anruflisten/*/data/*.json aller Remote-Branches sammeln."""
    def git(*a):
        return subprocess.run(["git", "-C", REPO, *a], capture_output=True, text=True, timeout=120)
    git("fetch", "-q", "origin", "+refs/heads/*:refs/remotes/origin/*")
    eigen = os.path.relpath(os.path.realpath(ausser), REPO) + "/" if ausser else None
    doms = set()
    for ref in git("for-each-ref", "--format=%(refname)", "refs/remotes").stdout.split():
        for pfad in git("ls-tree", "-r", "--name-only", ref, "--", "Anruflisten").stdout.splitlines():
            if not re.fullmatch(r"Anruflisten/[^/]+/data/[^/]+\.json", pfad) or (eigen and pfad.startswith(eigen)):
                continue
            try:
                daten = json.loads(git("show", f"{ref}:{pfad}").stdout)
            except ValueError:
                continue
            for r in daten if isinstance(daten, list) else []:
                if isinstance(r, dict) and domain(r.get("website")):
                    doms.add(domain(r.get("website")))
    return doms


# ---------- Excel ----------

def stufe(r, min_ma):
    n = ma_zahl(r["mitarbeiter"])
    if min_ma and n is not None and n < min_ma:
        return "C"
    if r["ansprechpartner"] != STRICH:
        return "A" if r["tel_ansprechpartner"] != STRICH else "B"
    return "C"


def prio_key(r):
    return (RANG.get(wort(r["eignung"]), 3), RANG.get(wort(r["signal_staerke"]), 3), r["stufe"],
            not eingeschraenkt(r["erreichbarkeit"]), r["bundesland"], r["ort"])


def laden(ordner, auftrag):
    ausschluss = alte_domains(ausser=ordner)
    rows, seen = [], set()
    for p in sorted(glob.glob(os.path.join(ordner, "data", "*.json"))):
        for r in json_liste(p) or []:
            extra = r.get("extra") if isinstance(r.get("extra"), dict) else {}
            quellen = als_liste(r.get("quellen") or r.get("mitarbeiter_quellen"))
            r = {k: leer(r.get(k)) for k in FELDER}
            r["extra"], r["quellen"] = extra, quellen
            if r["konkurrenz"] == STRICH:
                r["konkurrenz"] = "k. A."
            if r["firma"] == STRICH:
                continue
            k = domain(r["website"]) or re.sub(r"\W", "", r["firma"].lower())
            if k in ausschluss:
                print("schon früher geliefert:", r["firma"])
                continue
            if k in seen:
                print("Dublette:", r["firma"])
                continue
            seen.add(k)
            r["stufe"] = stufe(r, auftrag.get("min_mitarbeiter"))
            rows.append(r)
    rows.sort(key=prio_key)
    for i, r in enumerate(rows, 1):
        r["prio"] = i
    return rows


def kopf(ws, cols):
    for j, (name, w) in enumerate(cols, 1):
        c = ws.cell(1, j, name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = WRAP
        ws.column_dimensions[c.column_letter].width = w


def excel(ordner):
    auftrag = json.load(open(os.path.join(ordner, "auftrag.json"), encoding="utf-8"))
    titel = auftrag.get("titel", "Anrufliste")
    datum = auftrag.get("datum") or datetime.date.today().isoformat()
    datum_de = datetime.date.fromisoformat(datum).strftime("%d.%m.%Y")
    extra = []
    for e in auftrag.get("extra_spalten") or []:
        if isinstance(e, str):
            e = [e, e]
        if not isinstance(e, (list, tuple)) or not e:
            print("Zusatzspalte ignoriert:", e)
            continue
        key, lbl = str(e[0]), str(e[1] if len(e) > 1 else e[0])
        breite = e[2] if len(e) > 2 and isinstance(e[2], (int, float)) else 18
        extra.append((key, lbl, breite))
    signal_name = auftrag.get("signal_name", "Kaufsignal")
    verbund_name = auftrag.get("verbund_name", "Verbund/Gruppe")
    konk_name = auftrag.get("konkurrenz_name", "Anbieter schon da?")
    rows = laden(ordner, auftrag)

    wb = Workbook()
    ws = wb.active
    ws.title = "Anrufliste"
    cols = [("Prio", 6), (verbund_name, 16), ("Stufe", 7), ("Firma", 34), ("Ort", 18), ("Ansprechpartner", 24), ("Rolle", 26),
            ("Telefon Ansprechpartner", 22), ("Telefon Zentrale", 18), ("E-Mail", 24), ("Website", 28), ("Mitarbeiter", 14),
            ("Mitarbeiter sicher?", 11), ("Erreichbarkeit", 40), (signal_name, 40), ("Signal", 9), (konk_name, 22), ("Eignung", 28)]
    feste = {n for n, _ in cols} | {"Gesprächshinweis", "Angerufen am", "Status", "Ergebnis / Notiz"}
    cols += [(lbl if lbl not in feste else f"{lbl} (Zusatz)", w) for _, lbl, w in extra]
    cols += [("Gesprächshinweis", 38), ("Angerufen am", 13), ("Status", 18), ("Ergebnis / Notiz", 35)]
    kopf(ws, cols)
    sp = {name: j for j, (name, _) in enumerate(cols, 1)}
    for i, r in enumerate(rows, 2):
        ort = r["ort"] if r["bundesland"] == STRICH else f"{r['ort']} ({r['bundesland']})"
        vals = [r["prio"], r["verbund"], r["stufe"], r["firma"], ort, r["ansprechpartner"], r["rolle"], r["tel_ansprechpartner"],
                r["tel_zentrale"], r["email"], r["website"], r["mitarbeiter"], r["mitarbeiter_sicher"], r["erreichbarkeit"],
                r["kaufsignal"], r["signal_staerke"], r["konkurrenz"], r["eignung"]]
        vals += [leer(r["extra"].get(k)) for k, _, _ in extra]
        vals += [r["notiz"], None, None, None]
        for j, v in enumerate(vals, 1):
            ws.cell(i, j, v).alignment = WRAP
        ws.cell(i, sp["Stufe"]).fill = STUFE_FILL[r["stufe"]]
        if wort(r["eignung"]) == "hoch":
            ws.cell(i, sp["Eignung"]).fill = GRUEN
        if wort(r["signal_staerke"]) == "hoch":
            ws.cell(i, sp[signal_name]).fill = GRUEN
            ws.cell(i, sp["Signal"]).fill = GRUEN
        if r["konkurrenz"].lower().startswith("ja"):
            ws.cell(i, sp[konk_name]).fill = ROT
        if eingeschraenkt(r["erreichbarkeit"]):
            ws.cell(i, sp["Firma"]).fill = ORANGE
            ws.cell(i, sp["Firma"]).font = Font(bold=True)
            ws.cell(i, sp["Erreichbarkeit"]).fill = ORANGE
        if domain(r["website"]):
            url = r["website"].split(" ")[0]
            ws.cell(i, sp["Website"]).hyperlink = url if url.lower().startswith("http") else "https://" + url
            ws.cell(i, sp["Website"]).font = LINK
        if r["prio"] <= 15:
            for j in (sp["Prio"], sp["Firma"]):
                ws.cell(i, j).font = Font(bold=True)
    sc = ws.cell(1, sp["Status"]).column_letter
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{sc}2:{sc}{max(len(rows), 1) + 1}")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(cols)).column_letter}{len(rows) + 1}"

    q = wb.create_sheet("Mitarbeiter-Quellen")
    kopf(q, [("Firma", 34), ("Mitarbeiter", 14), ("sicher?", 9), ("Quellen", 90), ("Widerspruch / Hinweis", 60)])
    for i, r in enumerate(sorted(rows, key=lambda r: r["firma"].lower()), 2):
        for j, v in enumerate([r["firma"], r["mitarbeiter"], r["mitarbeiter_sicher"], "\n".join(map(str, r["quellen"])) or STRICH, r["widerspruch"]], 1):
            q.cell(i, j, v).alignment = WRAP
    q.freeze_panes = "B2"

    leitfaden = auftrag.get("leitfaden") or []
    if leitfaden:
        g = wb.create_sheet("Gesprächsleitfaden")
        g.column_dimensions["A"].width = 24
        g.column_dimensions["B"].width = 110
        g.cell(1, 1, f"{titel}: Gesprächsleitfaden (Stand {datum_de})").font = Font(bold=True)
        zeile = 3
        for block in leitfaden:
            if not isinstance(block, (list, tuple)) or len(block) < 2:
                continue
            g.cell(zeile, 1, leer(block[0])).font = Font(bold=True)
            eintraege = block[1] if isinstance(block[1], list) else [block[1]]
            for e in eintraege:
                g.cell(zeile, 2, leer(e)).alignment = WRAP
                zeile += 1
            zeile += 1

    lg = wb.create_sheet("Legende")
    lg.column_dimensions["A"].width = 22
    lg.column_dimensions["B"].width = 110
    n = {s: sum(r["stufe"] == s for r in rows) for s in "ABC"}
    min_ma = auftrag.get("min_mitarbeiter")
    legende = [
        (f"{titel}: Anrufliste, Stand {datum_de}", None), (None, None),
        ("Auftrag", f"Angebot: {auftrag.get('angebot', STRICH)}. Zielkunden: {auftrag.get('zielprofil', STRICH)}. Region: {auftrag.get('region', STRICH)}."),
        ("Prio", f"Reihenfolge zum Abtelefonieren: Eignung hoch vor mittel vor niedrig, dann {signal_name} stark vor schwach, dann Stufe A vor B vor C. Prio 1 bis 15 sind fett."),
        ("Stufe A", f"Entscheider mit Name UND eigener Durchwahl öffentlich ({n['A']} Firmen)"),
        ("Stufe B", f"Entscheider namentlich bekannt, nur Zentrale ({n['B']} Firmen)"),
        ("Stufe C", ("kein Entscheider gefunden" + (f" oder unter {min_ma} Mitarbeitern" if min_ma else "")) + f" ({n['C']} Firmen)"),
        (None, None),
        ("Eignung", "Wie gut die Firma zum Zielprofil passt (hoch grün), mit kurzer Begründung."),
        (signal_name, auftrag.get("signal_definition", "Konkreter Anlass, gerade jetzt anzurufen.") + " Spalte Signal: hoch (grün) / mittel / niedrig / – (nichts gefunden)."),
        (verbund_name, "Einkaufsverbund, Kooperation, Franchise oder Unternehmensgruppe, zu der die Firma gehört (– = keine gefunden)."),
        (konk_name, auftrag.get("konkurrenz_definition", "Nutzt die Firma schon einen Anbieter für das, was wir verkaufen?") + " ja = belegt (rot, Anbieter steht dabei), nein = Gegenteil belegt, k. A. = nichts gefunden."),
        ("Orange", "Firma schreibt, dass sie telefonisch eingeschränkt erreichbar ist (Telefonzeiten, Mittagspause, Freitag früh zu …). Zeiten stehen in der Spalte Erreichbarkeit."),
        ("Mitarbeiter", "Bestschätzung aus Website, Stellenanzeigen, LinkedIn, Presse, Registerauszügen. Alle Quellen im Blatt Mitarbeiter-Quellen."),
        ("Mitarbeiter sicher?", "hoch = eigene Angabe der Firma oder mehrere Quellen stimmen überein; mittel = eine plausible Quelle; niedrig = geschätzt oder widersprüchlich"),
        ("Gesprächshinweis", "Was beim Anruf hilft: Aufhänger, offene Stellen, neue Standorte, Besonderheiten."),
        ("Status", "Nach jedem Anruf auswählen: Termin; Rückruf vereinbart (Zeit in Notiz); Chef nicht erreicht (Durchwahl/andere Uhrzeit versuchen); kein Bedarf (6 Mon.) = später nochmal; gesperrt = nie wieder anrufen."),
        ("Strich", "– = Angabe nicht auffindbar"),
        ("Dubletten", "Firmen aus früheren Listen in ~/Anruflisten (und aus ~/Anruflisten/alt) sind herausgefiltert."),
    ]
    for i, (a, b) in enumerate(legende, 1):
        lg.cell(i, 1, a).alignment = WRAP
        lg.cell(i, 2, b).alignment = WRAP
        if a and b:
            lg.cell(i, 1).font = Font(bold=True)
    lg.cell(1, 1).font = Font(bold=True)

    path = os.path.join(ordner, f"{slug(titel)}-Anrufliste-{datum}.xlsx")
    # Recherchierte Texte nie als Formel speichern (z. B. ein Firmenname "=HYPERLINK(...)")
    for blatt in wb.worksheets:
        for zeile in blatt.iter_rows():
            for c in zeile:
                if c.data_type == "f":
                    c.data_type = "s"
    wb.save(path)
    hoch = sum(wort(r["signal_staerke"]) == "hoch" for r in rows)
    print(f"{path}\n{len(rows)} Firmen, Stufe A {n['A']} / B {n['B']} / C {n['C']}, {signal_name} hoch {hoch}, "
          f"Eignung hoch {sum(wort(r['eignung']) == 'hoch' for r in rows)}, eingeschränkt erreichbar {sum(eingeschraenkt(r['erreichbarkeit']) for r in rows)}")


def neu(titel):
    datum = datetime.date.today().isoformat()
    # im Repo entstehen Listen auf verschiedenen Branches: Uhrzeit verhindert gleiche Ordnernamen
    zeit = datetime.datetime.now().strftime("-%H%M") if REPO else ""
    basis = os.path.join(ROOT, f"{datum}{zeit}-{slug(titel)}")
    ordner, i = basis, 2
    while os.path.exists(ordner):
        ordner, i = f"{basis}-{i}", i + 1
    os.makedirs(os.path.join(ordner, "data"))
    os.makedirs(os.path.join(ROOT, "alt"), exist_ok=True)
    print(ordner)


def ausschluss(ordner):
    doms = sorted(alte_domains(ausser=ordner))
    open(os.path.join(ordner, "ausschluss_domains.txt"), "w", encoding="utf-8").write("\n".join(doms) + ("\n" if doms else ""))
    print(f"{len(doms)} früher gelieferte Firmen in {os.path.join(ordner, 'ausschluss_domains.txt')}")


if __name__ == "__main__":
    if len(sys.argv) != 3 or sys.argv[1] not in ("neu", "ausschluss", "excel"):
        sys.exit(__doc__)
    {"neu": neu, "ausschluss": ausschluss, "excel": excel}[sys.argv[1]](os.path.expanduser(sys.argv[2]) if sys.argv[1] != "neu" else sys.argv[2])
